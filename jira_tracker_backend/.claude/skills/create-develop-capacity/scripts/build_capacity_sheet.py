#!/usr/bin/env python3
"""Generate the Rolling Sprint Capacity workbook from a JSON config.

Usage:
    python build_capacity_sheet.py <config.json>

The workbook encodes the team's rolling-capacity rule:
    each sprint's capacity = ROUND(AVERAGE(previous sprint capacity,
                                           previous sprint actual), 0)
Excel's ROUND rounds halves away from zero, so 4.5 -> 5 (as the team wants).

Config schema (all keys optional except `developers`):
{
  "output": "~/Desktop/developer_capacity_weekly.xlsx",
  "title":  "Rolling Sprint Capacity — Developers (tickets / sprint)",
  "sprints": [
    {"name": "SPRINT 1", "dates": "Jun 29 – Jul 3",  "current": false},
    {"name": "SPRINT 2", "dates": "Jul 6 – Jul 10",  "current": false},
    {"name": "SPRINT 3", "dates": "Jul 13 – Jul 17", "current": true },
    {"name": "SPRINT 4", "dates": "projected",       "current": false}
  ],
  "developers": [
    {"name": "Sidney", "capacity": 5,  "actual": 4},
    {"name": "Muscan", "capacity": 10, "actual": null}
  ]
}

Notes
- Only Sprint 1 capacity/actual are entered per developer; Sprints 2-4 are
  computed (rolling). Later sprints' actuals stay blank for manual entry —
  the formulas pick them up automatically once typed in.
- `actual: null` (or omitted) leaves that developer's actual blank.
"""
import json
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side

DARK = "1F3864"
GREY = "666666"
GREEN = "2E7D32"

DEFAULT_SPRINTS = [
    {"name": "SPRINT 1", "dates": "", "current": False},
    {"name": "SPRINT 2", "dates": "", "current": False},
    {"name": "SPRINT 3", "dates": "", "current": False},
    {"name": "SPRINT 4", "dates": "projected", "current": False},
]


def _col(idx):
    """1-based column index -> letter (A, B, ...)."""
    return openpyxl.utils.get_column_letter(idx)


def build(config):
    title = config.get(
        "title", "Rolling Sprint Capacity — Developers (tickets / sprint)"
    )
    sprints = config.get("sprints", DEFAULT_SPRINTS)
    devs = config["developers"]

    # Per-sprint column layout: sprint i -> (capacity_col, actual_col).
    # Sprints 1..n-1 have Capacity+Actual; the final sprint is projected
    # capacity only. Column B is the first data column.
    layout = []
    c = 2
    for i, _ in enumerate(sprints):
        last_sprint = i == len(sprints) - 1
        cap = c
        act = None if last_sprint else c + 1
        layout.append((cap, act))
        c += 1 if last_sprint else 2
    n_cols = c - 1  # last used column index

    light = Side(style="thin", color="DDDDDD")
    med = Side(style="medium", color=DARK)
    acc = Side(style="medium", color=GREEN)
    box = Border(left=light, right=light, top=light, bottom=light)
    ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rolling Sprint Capacity"
    ws.sheet_view.showGridLines = False

    last_letter = _col(n_cols)
    for i in range(1, n_cols + 1):
        ws.column_dimensions[_col(i)].width = 22 if i == 1 else 11
    ws.column_dimensions[_col(n_cols)].width = 13

    # Row 1 — title
    ws.merge_cells(f"A1:{last_letter}1")
    cell = ws["A1"]
    cell.value = title
    cell.font = Font(bold=True, size=15, color=DARK)
    cell.alignment = left
    ws.row_dimensions[1].height = 26

    # Row 2 — rule
    ws.merge_cells(f"A2:{last_letter}2")
    cell = ws["A2"]
    cell.value = (
        "Rule: each sprint's capacity = average(previous sprint capacity, "
        "previous sprint actual). Capacity is computed automatically · enter "
        "actuals in the Actual columns."
    )
    cell.font = Font(italic=True, size=10, color=GREY)
    cell.alignment = left
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 6

    # Row 4 — sprint headers, Row 5 — date ranges
    for (cap, act), sp in zip(layout, sprints):
        cur = sp.get("current", False)
        color = GREEN if cur else DARK
        end = act if act else cap
        ws.merge_cells(f"{_col(cap)}4:{_col(end)}4")
        hc = ws[f"{_col(cap)}4"]
        hc.value = (sp["name"] + "  ● CURRENT") if cur else sp["name"]
        hc.font = Font(bold=True, size=11, color=color)
        hc.alignment = ctr
        hc.border = Border(bottom=Side(style="thin", color=color))
        ws.merge_cells(f"{_col(cap)}5:{_col(end)}5")
        dc = ws[f"{_col(cap)}5"]
        dc.value = sp.get("dates", "")
        dc.font = Font(italic=True, size=9, bold=cur, color=color if cur else GREY)
        dc.alignment = ctr
    ws["A4"].border = Border(bottom=Side(style="thin", color=DARK))
    ws.row_dimensions[4].height = 18
    ws.row_dimensions[5].height = 15

    # Row 6 — column headers
    ws.cell(row=6, column=1, value="Developer").font = Font(
        bold=True, size=10, color=DARK
    )
    ws.cell(row=6, column=1).alignment = ctr
    ws.cell(row=6, column=1).border = Border(
        bottom=med, top=Side(style="thin", color=DARK)
    )
    for i, (cap, act) in enumerate(layout):
        last_sprint = i == len(layout) - 1
        cur = sprints[i].get("current", False)
        hcolor = GREEN if cur else DARK
        cap_label = "Capacity\n(proj.)" if last_sprint else "Capacity"
        cell = ws.cell(row=6, column=cap, value=cap_label)
        cell.font = Font(bold=True, size=10, color=hcolor)
        cell.alignment = ctr
        cell.border = Border(bottom=med, top=Side(style="thin", color=DARK))
        if act:
            cell = ws.cell(row=6, column=act, value="Actual")
            cell.font = Font(bold=True, size=10, color=hcolor)
            cell.alignment = ctr
            cell.border = Border(bottom=med, top=Side(style="thin", color=DARK))
    ws.row_dimensions[6].height = 26

    # Developer rows
    first = 7
    last = first + len(devs) - 1
    cap0, act0 = layout[0]  # Sprint 1 capacity/actual columns
    for idx, dev in enumerate(devs):
        r = first + idx
        ws.cell(row=r, column=1, value=dev["name"]).font = Font(bold=True, size=10)
        ws.cell(row=r, column=cap0, value=dev.get("capacity"))
        actual = dev.get("actual")
        if actual is not None:
            ws.cell(row=r, column=act0, value=actual)
        # Rolling capacity for sprints 2..n
        prev_cap = cap0
        prev_act = act0
        for cap, act in layout[1:]:
            pc, pa = _col(prev_cap), _col(prev_act)
            if act is None:  # projected final sprint — no actual of its own
                formula = (
                    f'=IF({pa}{r}="",{pc}{r},ROUND(AVERAGE({pc}{r},{pa}{r}),0))'
                )
            else:
                formula = (
                    f'=IF({pa}{r}="",{pc}{r},ROUND(AVERAGE({pc}{r},{pa}{r}),0))'
                    if prev_cap != cap0
                    else f"=ROUND(AVERAGE({pc}{r},{pa}{r}),0)"
                )
            ws.cell(row=r, column=cap, value=formula)
            prev_cap, prev_act = cap, act
        for col in range(1, n_cols + 1):
            cell = ws.cell(row=r, column=col)
            cell.border = box
            cell.alignment = left if col == 1 else ctr
            if col > 1:
                cell.number_format = "0"
        ws.row_dimensions[r].height = 18

    # TOTAL row
    tr = last + 1
    ws.cell(row=tr, column=1, value="TOTAL").font = Font(bold=True, size=10, color=DARK)
    ws.cell(row=tr, column=1).alignment = left
    for col in range(2, n_cols + 1):
        letter = _col(col)
        cell = ws.cell(row=tr, column=col, value=f"=SUM({letter}{first}:{letter}{last})")
        cell.font = Font(bold=True, size=10, color=DARK)
        cell.alignment = ctr
        cell.number_format = "0"
    for col in range(1, n_cols + 1):
        ws.cell(row=tr, column=col).border = Border(top=med, bottom=med)
    ws.row_dimensions[tr].height = 20

    # Green accent border down the current sprint's column(s)
    for i, (cap, act) in enumerate(layout):
        if not sprints[i].get("current", False):
            continue
        cols = [cap] + ([act] if act else [])
        for rr in range(4, tr + 1):
            for j, col in enumerate(cols):
                cell = ws.cell(row=rr, column=col)
                b = cell.border
                ws.cell(row=rr, column=col).border = Border(
                    left=acc if j == 0 else b.left,
                    right=acc if j == len(cols) - 1 else b.right,
                    top=acc if rr == 4 else b.top,
                    bottom=acc if rr == tr else b.bottom,
                )

    # SUMMARY section
    cap1 = _col(layout[0][0])
    act1 = _col(layout[0][1])
    cap2 = _col(layout[1][0]) if len(layout) > 1 else cap1
    # last non-projected capacity column = the current/latest real sprint
    latest_cap = _col(layout[-2][0]) if len(layout) > 2 else cap2
    sr = tr + 2
    ws.merge_cells(f"A{sr}:{last_letter}{sr}")
    cell = ws[f"A{sr}"]
    cell.value = "SUMMARY"
    cell.font = Font(bold=True, size=11, color=DARK)
    cell.alignment = left
    cell.border = Border(bottom=Side(style="thin", color=DARK))
    ws.row_dimensions[sr].height = 20

    rng = f"{first}:%s{last}"
    summary = [
        ("Team total capacity — Sprint 1", f"=SUM({cap1}{first}:{cap1}{last})", "0"),
        ("Team total actual — Sprint 1", f"=SUM({act1}{first}:{act1}{last})", "0"),
        ("Team total capacity — Sprint 2", f"=SUM({cap2}{first}:{cap2}{last})", "0"),
        (
            "Team total capacity — latest sprint",
            f"=SUM({latest_cap}{first}:{latest_cap}{last})",
            "0",
        ),
        (
            "Avg capacity / dev — latest sprint",
            f"=AVERAGE({latest_cap}{first}:{latest_cap}{last})",
            "0.0",
        ),
        (
            "Sprint 1 utilization (actual / plan)",
            f"=SUM({act1}{first}:{act1}{last})/SUM({cap1}{first}:{cap1}{last})",
            "0.0%",
        ),
    ]
    row = sr + 1
    mid = _col(min(5, n_cols))
    val_start = _col(min(6, n_cols))
    for label, formula, fmt in summary:
        ws.merge_cells(f"A{row}:{mid}{row}")
        lc = ws[f"A{row}"]
        lc.value = label
        lc.font = Font(size=10)
        lc.alignment = left
        for col in range(1, min(5, n_cols) + 1):
            ws.cell(row=row, column=col).border = box
        ws.merge_cells(f"{val_start}{row}:{last_letter}{row}")
        vc = ws[f"{val_start}{row}"]
        vc.value = formula
        vc.font = Font(bold=True, size=10, color=DARK)
        vc.alignment = ctr
        vc.number_format = fmt
        for col in range(min(6, n_cols), n_cols + 1):
            ws.cell(row=row, column=col).border = box
        ws.row_dimensions[row].height = 18
        row += 1

    out = os.path.expanduser(
        config.get("output", "~/Desktop/developer_capacity_weekly.xlsx")
    )
    wb.save(out)
    return out


def main():
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(1)
    with open(sys.argv[1], encoding="utf-8") as fh:
        config = json.load(fh)
    out = build(config)
    print(f"Saved: {out}")


if __name__ == "__main__":
    main()
