#!/usr/bin/env python3
"""Product Status Report — build one prefilled report per product, and check a filled one.

`build` consumes a worklog-tracker Timesheet and emits one Product Status Report per product
section it finds, prefilling only what the timesheet can honestly source (§3's roster, hours and
cross-product splits) plus whatever last week's report can carry forward. Every cell with no
source is left an explicit TODO rather than guessed.

`check` re-reads a report a human has filled in and reports what would embarrass the sender:
leftover template EXAMPLE rows, an Hours/Day divisor that contradicts the stated week,
implausible days, missing cross-product split notes, and internal contradictions.
"""

import argparse
import atexit
import copy
import datetime as dt
import functools
import re
import shutil
import subprocess
import sys
import tempfile
from collections import namedtuple
from pathlib import Path

import openpyxl

# --- The template's cell map. Row numbers are the bundled product_status_template.xlsx, and
# --- these constants are the single source of truth for them: nothing else hardcodes a row.
CELL_PRODUCT, CELL_REPORTER, CELL_WEEK, CELL_SUBMITTED = "C3", "C4", "C5", "C6"
S1_ROWS = (10, 11)                    # §1 what this product is
S2_ROWS = (15, 19)                    # §2 what we are building right now
S3_FIRST, S3_LAST, S3_TOTAL = 24, 36, 37
S4_LABELS = {40: "start", 41: "enterprise_deploy", 42: "next_milestone", 43: "enterprise_today"}
S5_GOLIVE, S5_FIRST, S5_LAST = 46, 49, 64
S6_FIRST, S6_LAST = 68, 77            # §6 blockers — never carried, always starts empty
S7_ROWS = (80, 83)

COL_NAME, COL_ROLE, COL_HOURS, COL_PER_DAY, COL_NOTES = 2, 3, 4, 5, 6
TABLE_COLUMNS = range(COL_NAME, COL_NOTES + 1)

TODO = "TODO — needs input"
CARRY_SUFFIX = "  [carried — confirm]"
EXAMPLE_MARKER = "EXAMPLE (delete this row)"
IMPLAUSIBLE_HOURS_PER_DAY = 10.0
TEMPLATE = Path(__file__).with_name("product_status_template.xlsx")

# A window opening after midday contributes no working day — the worklog's Thursday-18:00
# anchor is an evening cutoff, not a day of work.
EVENING_CUTOFF_HOUR = 12

# --- The report filename is written in one place and parsed in another; both come from here,
# --- so the naming scheme cannot drift between the writer and its readers.
REPORT_PREFIX = "Product_Status_Report"
SPAN_PATTERN = r"\d{4}-\d{2}-\d{2}_\d{4}-\d{2}-\d{2}"
REPORT_RE = re.compile(rf"^{REPORT_PREFIX}_(?P<product>.+?)_(?P<span>{SPAN_PATTERN})$")
READABLE_SUFFIXES = (".xlsx", ".numbers")

ReportSpec = namedtuple("ReportSpec", "product week days span")


def die(message):
    """Abort with a one-line reason — every failure here is the operator's to fix."""
    sys.exit(f"error: {message}")


def report_filename(product, span):
    """The canonical filename for one product's report over one window."""
    return f"{REPORT_PREFIX}_{product}_{span}.xlsx"


def report_glob(product="*", span="*"):
    """A glob matching a product's reports in any readable format."""
    return f"{REPORT_PREFIX}_{product}_{span}"


def report_identity(path):
    """The (product, span) a report filename encodes, or None if it is not one of ours."""
    match = REPORT_RE.match(Path(path).stem)
    return (match.group("product"), match.group("span")) if match else None


# --------------------------------------------------------------------------- reading input

@functools.lru_cache(maxsize=1)
def conversion_dir():
    """One temp directory per process for .numbers conversions, removed at exit.

    Deliberately not a `with` block: the converted files must outlive this call and stay
    readable for the whole run, so the directory is reclaimed at exit instead.
    """
    path = Path(tempfile.mkdtemp(prefix="product-report-"))
    atexit.register(shutil.rmtree, path, ignore_errors=True)
    return path


@functools.lru_cache(maxsize=None)
def convert_numbers(path):
    """Export an Apple .numbers file to .xlsx through Numbers.app, once per process."""
    source = Path(path)
    out = conversion_dir() / f"{source.stem}.xlsx"
    script = (f'tell application "Numbers"\n'
              f'  set d to open POSIX file "{source}"\n'
              f'  export d to (POSIX file "{out}") as Microsoft Excel\n'
              f'  close d saving no\n'
              f'end tell')
    result = subprocess.run(["osascript", "-e", script],
                            capture_output=True, text=True, check=False)
    if result.returncode != 0 or not out.exists():
        die(f"could not convert {source.name} via Numbers.app: {result.stderr.strip()}")
    return out


def as_xlsx(path):
    """An .xlsx path for `path` — a .numbers is a file a human resaved, so convert it."""
    path = Path(path).expanduser()
    if path.suffix.lower() != ".numbers":
        return path
    return convert_numbers(str(path.resolve()))


def load_sheet(path, name=None):
    """The worksheet to read values from — the one place load options are decided."""
    book = openpyxl.load_workbook(as_xlsx(path), data_only=True)
    return book[name] if name and name in book.sheetnames else book.active


def parse_timesheet(path):
    """Return (window_text, {product: [(developer, hours, title), ...]}) from a Timesheet."""
    sheet = load_sheet(path, "Timesheet")
    window = str(sheet["A1"].value or "")

    # Everything above the `Developer | Hours | Title` header is title and caveat prose; a
    # single-cell row up there looks exactly like a shaded section header if you let it.
    header_row = next((cell.row for cell in sheet["A"]
                       if str(cell.value or "").strip() == "Developer"), None)
    if header_row is None:
        die(f"no 'Developer' header row in {Path(path).name} — is this a worklog Timesheet?")

    products, current = {}, None
    for row in sheet.iter_rows(min_row=header_row + 1, max_col=3):
        name, hours, title = (cell.value for cell in row)
        name = str(name).strip() if name is not None else ""
        if not name or name == "TOTAL":
            continue
        if hours is None and not title:          # a shaded section header, not a work item
            current = name
            products.setdefault(current, [])
            continue
        if current and isinstance(hours, (int, float)):
            products[current].append((name, float(hours), str(title or "").strip()))
    if not products:
        die(f"no product sections found in {Path(path).name} — is this a worklog Timesheet?")
    return window, products


def window_dates(window_text):
    """The (start, end) datetimes named in the timesheet's A1 window string."""
    stamps = re.findall(r"(\d{4}-\d{2}-\d{2})(?:[ T](\d{2}):(\d{2}))?", window_text)
    if len(stamps) < 2:
        return None, None

    def parse(stamp):
        day, hour, minute = stamp
        return dt.datetime.strptime(day, "%Y-%m-%d").replace(
            hour=int(hour or 0), minute=int(minute or 0))

    return parse(stamps[0]), parse(stamps[-1])


def working_days(start, end):
    """Mon–Fri days in [start, end], skipping a start day whose window opens in the evening."""
    if not start or not end:
        return 5
    opens_late = start.hour >= EVENING_CUTOFF_HOUR
    first = start.date() + dt.timedelta(days=1) if opens_late else start.date()
    days = [first + dt.timedelta(days=n) for n in range((end.date() - first).days + 1)]
    return sum(1 for day in days if day.weekday() < 5) or 5


def week_label(start, end):
    """The Reporting Week text, in the template's own date style."""
    if not start or not end:
        return TODO
    if start.year != end.year:
        return f"{start:%B %-d, %Y} - {end:%B %-d, %Y}"
    return f"{start:%B %-d} - {end:%B %-d, %Y}"


def window_span(start, end):
    """The filename span for a window."""
    if not start or not end:
        return f"{dt.date.today():%Y-%m-%d}_{dt.date.today():%Y-%m-%d}"
    return f"{start:%Y-%m-%d}_{end:%Y-%m-%d}"


# --------------------------------------------------------------------------- carry-forward

def empty_carry():
    """The carry payload for a product with no previous report."""
    return {"reporter": "", "section1": [], "section4": {}, "golive": "",
            "section7": [], "roles": {}, "section5": []}


def find_carry_source(out_dir, product, exclude):
    """The most recent previous report for this product — the artifact IS the store."""
    candidates = [path
                  for suffix in READABLE_SUFFIXES
                  for path in Path(out_dir).expanduser().glob(report_glob(product) + suffix)
                  if path.resolve() != Path(exclude).resolve()]
    if not candidates:
        return None
    return max(candidates, key=lambda path: path.stat().st_mtime)


def read_carry(path):
    """Pull the slow-moving answers out of a previous report. §6 blockers are never carried."""
    sheet = load_sheet(path)

    def text(coord):
        value = sheet[coord].value
        return str(value).replace(CARRY_SUFFIX.strip(), "").strip() if value else ""

    carried = {
        "reporter": text(CELL_REPORTER),
        "section1": [text(f"B{row}") for row in range(S1_ROWS[0], S1_ROWS[1] + 1)],
        "section4": {label: text(f"C{row}") for row, label in S4_LABELS.items()},
        "golive": text(f"C{S5_GOLIVE}"),
        "section7": [text(f"B{row}") for row in range(S7_ROWS[0], S7_ROWS[1] + 1)],
        "roles": {},
        "section5": [],
    }
    for row in range(S3_FIRST, S3_LAST + 1):
        name, role = sheet.cell(row, COL_NAME).value, sheet.cell(row, COL_ROLE).value
        if name and role:
            carried["roles"][str(name).strip()] = str(role).strip()
    for row in range(S5_FIRST, S5_LAST + 1):
        cells = [sheet.cell(row, col).value for col in TABLE_COLUMNS]
        if any(cells) and EXAMPLE_MARKER not in str(cells[0] or ""):
            carried["section5"].append([str(cell).strip() if cell else "" for cell in cells])
    return carried


# --------------------------------------------------------------------------- build

def developer_rows(items, other_products):
    """Collapse a product's work items into one §3 row per developer, hours descending."""
    totals, counts = {}, {}
    for name, hours, _title in items:
        totals[name] = totals.get(name, 0.0) + hours
        counts[name] = counts.get(name, 0) + 1

    rows = []
    for name in sorted(totals, key=lambda key: -totals[key]):
        notes = [f"Estimate from commit volume ({counts[name]} items)"]
        elsewhere = [f"{product} ({hours:.1f} hrs)"
                     for product, hours in sorted(other_products.get(name, {}).items())]
        if elsewhere:
            notes.insert(0, "Split with " + ", ".join(elsewhere))
        rows.append((name, round(totals[name], 1), " - ".join(notes)))
    return rows


def cross_product_hours(products):
    """{developer: {product: hours}} — the source of every 'Split with …' note."""
    totals = {}
    for product, items in products.items():
        for name, hours, _title in items:
            totals.setdefault(name, {}).setdefault(product, 0.0)
            totals[name][product] += hours
    return {name: {product: round(hours, 1) for product, hours in by_product.items()}
            for name, by_product in totals.items()}


def clear_rows(sheet, first, last, columns=TABLE_COLUMNS):
    """Blank a block of table cells."""
    for row in range(first, last + 1):
        for col in columns:
            sheet.cell(row, col).value = None


def write_section3(sheet, rows, days, roles):
    """Fill the roster: one row per developer, with Hours/Day as a live formula."""
    body_font = copy.copy(sheet.cell(S3_FIRST + 1, COL_NAME).font)   # the normal 10pt data rows
    clear_rows(sheet, S3_FIRST, S3_LAST)
    for offset, (name, hours, notes) in enumerate(rows):
        row = S3_FIRST + offset
        if row > S3_LAST:
            print(f"warning: {len(rows)} developers exceeds the template's "
                  f"{S3_LAST - S3_FIRST + 1} rows — extras dropped", file=sys.stderr)
            break
        values = {COL_NAME: name, COL_ROLE: roles.get(name, ""), COL_HOURS: hours,
                  COL_PER_DAY: f"=D{row}/{days}", COL_NOTES: notes}
        for col, value in values.items():
            cell = sheet.cell(row, col)
            cell.value = value
            cell.font = body_font
            cell.number_format = {COL_HOURS: "0.00", COL_PER_DAY: "0.0"}.get(col, "@")


def mark_carried(value):
    """A carried answer, visibly flagged — or the TODO that says nothing was carried."""
    return f"{value}{CARRY_SUFFIX}" if value else TODO


def write_carried_paragraph(sheet, rows, texts):
    """Fill a multi-row free-text section: the first row always answers, later rows may be blank."""
    for offset, row in enumerate(range(rows[0], rows[1] + 1)):
        text = texts[offset] if offset < len(texts) else ""
        if offset == 0:
            sheet[f"B{row}"] = mark_carried(text)
            continue
        sheet[f"B{row}"] = f"{text}{CARRY_SUFFIX}" if text else None


def write_carried(sheet, carried):
    """Fill the slow-moving sections, each visibly marked so nothing stale ships unnoticed."""
    sheet[CELL_REPORTER] = mark_carried(carried["reporter"])
    write_carried_paragraph(sheet, S1_ROWS, carried["section1"])
    write_carried_paragraph(sheet, S7_ROWS, carried["section7"])
    for row, label in S4_LABELS.items():
        sheet[f"C{row}"] = mark_carried(carried["section4"].get(label, ""))
    sheet[f"C{S5_GOLIVE}"] = mark_carried(carried["golive"])
    for offset, values in enumerate(carried["section5"][:S5_LAST - S5_FIRST + 1]):
        for col, value in zip(TABLE_COLUMNS, values):
            sheet.cell(S5_FIRST + offset, col).value = value or None
        sheet.cell(S5_FIRST + offset, COL_NAME).value = mark_carried(values[0])


def blank_unsourced(sheet):
    """§2 and §6 have no source and are never carried — leave them explicitly unanswered."""
    for row in range(S2_ROWS[0], S2_ROWS[1] + 1):
        sheet[f"B{row}"] = TODO if row == S2_ROWS[0] else None
    clear_rows(sheet, S6_FIRST, S6_LAST)


def build_report(spec, rows, out_path, carried):
    """Write one product's report from the template."""
    shutil.copyfile(TEMPLATE, out_path)
    book = openpyxl.load_workbook(out_path)
    sheet = book.active

    sheet[CELL_PRODUCT] = spec.product
    sheet[CELL_WEEK] = spec.week
    sheet[CELL_SUBMITTED] = f"{dt.date.today():%B %-d, %Y}"
    clear_rows(sheet, S5_FIRST, S5_LAST)                    # drop the template's EXAMPLE row
    write_section3(sheet, rows, spec.days, carried["roles"])
    write_carried(sheet, carried)
    blank_unsourced(sheet)

    book.save(out_path)
    return out_path


# --------------------------------------------------------------------------- check

def check_report(path):
    """Everything that would embarrass the sender, as (severity, message) pairs."""
    sheet = load_sheet(path)
    findings = []
    for row in sheet.iter_rows():
        for cell in row:
            value = str(cell.value or "")
            if EXAMPLE_MARKER in value:
                findings.append(("ERROR", f"{cell.coordinate}: template EXAMPLE row still "
                                          f"present — delete it"))
            elif TODO in value:
                findings.append(("ERROR", f"{cell.coordinate}: still reads '{TODO}'"))
            elif CARRY_SUFFIX.strip() in value:
                findings.append(("WARN", f"{cell.coordinate}: carried from last week and not "
                                         f"yet confirmed"))
    return findings + check_section3(sheet) + check_consistency(sheet) + \
        check_cross_product(path, sheet)


def check_section3(sheet):
    """The roster's own arithmetic: a shared divisor, plausible days, and a Role per person."""
    findings, divisors = [], set()
    for row in range(S3_FIRST, S3_LAST + 1):
        name = sheet.cell(row, COL_NAME).value
        hours, per_day = sheet.cell(row, COL_HOURS).value, sheet.cell(row, COL_PER_DAY).value
        if not name:
            continue
        if not sheet.cell(row, COL_ROLE).value:
            findings.append(("WARN", f"B{row}: {name} has no Role"))
        if not isinstance(hours, (int, float)) or not isinstance(per_day, (int, float)):
            continue
        if per_day > IMPLAUSIBLE_HOURS_PER_DAY:
            findings.append(("WARN", f"E{row}: {name} at {per_day:.1f} h/day — above "
                                     f"{IMPLAUSIBLE_HOURS_PER_DAY:g}, verify or explain it"))
        if per_day:
            divisors.add(round(hours / per_day))
    if len(divisors) > 1:
        findings.append(("ERROR", f"§3 Hours/Day uses inconsistent divisors {sorted(divisors)} "
                                  f"— every row must divide by the same working-day count"))
    return findings


def row_for(label):
    """The template row carrying one of §4's four timeline answers."""
    return next(row for row, name in S4_LABELS.items() if name == label)


def check_consistency(sheet):
    """Contradictions between sections, and dates a reader cannot act on."""
    findings = []
    enterprise = str(sheet[f"C{row_for('enterprise_today')}"].value or "").strip().lower()
    undeployed = " ".join([
        str(sheet[f"C{row_for('enterprise_deploy')}"].value or ""),
        str(sheet[f"C{S5_GOLIVE}"].value or ""),
        *(str(sheet[f"B{row}"].value or "") for row in range(S2_ROWS[0], S2_ROWS[1] + 1)),
    ]).lower()
    unshipped = [p for p in ("tbd", "to be determined", "phase 0") if p in undeployed]
    if enterprise.startswith("yes") and unshipped:
        findings.append(("WARN", f"§4 claims enterprise-level 'yes' while §2/§4/§5 read "
                                 f"pre-production ({', '.join(unshipped)}) — these contradict"))

    submitted = str(sheet[CELL_SUBMITTED].value or "")
    if submitted and f"{dt.date.today():%-d, %Y}" not in submitted:
        findings.append(("WARN", f"{CELL_SUBMITTED}: Date Submitted reads '{submitted}', "
                                 f"today is {dt.date.today():%B %-d, %Y}"))

    month = (r"(January|February|March|April|May|June|July|August|September|October|November"
             r"|December)")
    for coord in [f"C{row}" for row in S4_LABELS] + [f"C{S5_GOLIVE}"]:
        value = str(sheet[coord].value or "")
        if re.search(rf"{month}\s+\d{{1,2}}\b", value) and not re.search(r"\b\d{4}\b", value):
            findings.append(("WARN", f"{coord}: '{value}' names a date with no year"))
        if re.search(r"\+\s*\d+\s*(day|week|month)", value, re.IGNORECASE):
            findings.append(("WARN", f"{coord}: '{value}' expresses a date as arithmetic — "
                                     f"write the actual date a reader can put in a calendar"))
    return findings


def section3_names(sheet):
    """Every developer named in a report's roster."""
    return {str(sheet.cell(row, COL_NAME).value).strip()
            for row in range(S3_FIRST, S3_LAST + 1) if sheet.cell(row, COL_NAME).value}


def sibling_rosters(path, product, span):
    """{other product: its roster} for the same window, read off the other reports on disk."""
    rosters = {}
    for other in Path(path).parent.glob(report_glob(span=span) + ".*"):
        identity = report_identity(other)
        if not identity or other.suffix.lower() not in READABLE_SUFFIXES:
            continue
        if identity[0] == product or other.resolve() == Path(path).resolve():
            continue
        rosters[identity[0]] = section3_names(load_sheet(other))
    return rosters


def check_cross_product(path, sheet):
    """A developer on two products' reports for the same window must say so in Notes.

    §3 asks for hours on THIS product only, so an unsplit row silently claims a developer's
    whole week for one product — which is how a 39.5h SETHI row hid 10.5h of NEXUS work.
    """
    identity = report_identity(path)
    if not identity:
        return []
    rosters = sibling_rosters(path, *identity)

    findings = []
    for row in range(S3_FIRST, S3_LAST + 1):
        name = sheet.cell(row, COL_NAME).value
        if not name:
            continue
        elsewhere = [other for other, names in rosters.items() if str(name).strip() in names]
        if elsewhere and "split" not in str(sheet.cell(row, COL_NOTES).value or "").lower():
            findings.append(("WARN", f"F{row}: {name} also appears on the "
                                     f"{', '.join(elsewhere)} report for this window, "
                                     f"but Notes names no split"))
    return findings


# --------------------------------------------------------------------------- cli

def report_one(spec, items, elsewhere, out_dir, no_carry):
    """Build one product's report and print what it drew on."""
    out_path = Path(out_dir) / report_filename(spec.product, spec.span)
    others = {name: {p: h for p, h in by.items() if p != spec.product}
              for name, by in elsewhere.items()}
    source = None if no_carry else find_carry_source(out_dir, spec.product, out_path)
    carried = read_carry(source) if source else empty_carry()
    rows = developer_rows(items, others)
    build_report(spec, rows, out_path, carried)

    print(f"{spec.product}: {out_path}")
    print(f"  week={spec.week}  working_days={spec.days}  developers={len(rows)}  "
          f"total={sum(row[1] for row in rows):.1f}h")
    print(f"  carried from: {source.name if source else 'nothing — first report for this product'}")


def run_build(args):
    """Emit one report per product section in the timesheet."""
    window, products = parse_timesheet(args.timesheet)
    start, end = window_dates(window)
    week = args.week or week_label(start, end)
    days = args.working_days or working_days(start, end)
    span = window_span(start, end)

    out_dir = Path(args.out_dir).expanduser()
    out_dir.mkdir(parents=True, exist_ok=True)
    wanted = [p for p in products if not args.product or p.upper() == args.product.upper()]
    if not wanted:
        die(f"product {args.product!r} not in the timesheet — found {sorted(products)}")

    elsewhere = cross_product_hours(products)
    for product in wanted:
        report_one(ReportSpec(product, week, days, span),
                   products[product], elsewhere, out_dir, args.no_carry)


def run_check(args):
    """Audit a filled report; exit non-zero when anything blocking remains."""
    findings = check_report(args.report)
    for severity, message in findings:
        print(f"{severity}: {message}")
    errors = sum(1 for severity, _ in findings if severity == "ERROR")
    print(f"\n{len(findings)} finding(s), {errors} blocking.")
    return 1 if errors else 0


def main():
    """Parse the command line and dispatch."""
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)

    build = sub.add_parser("build", help="prefill one Product Status Report per product")
    build.add_argument("--timesheet", required=True,
                       help="worklog-tracker Timesheet .xlsx (or .numbers, if hand-corrected)")
    build.add_argument("--out-dir", default="~/Desktop")
    build.add_argument("--product", help="only this product section (default: every one found)")
    build.add_argument("--week", help="override the Reporting Week text")
    build.add_argument("--working-days", type=int, help="override the Hours/Day divisor")
    build.add_argument("--no-carry", action="store_true",
                       help="do not carry last week's answers forward")

    check = sub.add_parser("check", help="audit a filled-in report before it is sent")
    check.add_argument("report", help="a filled Product Status Report .xlsx/.numbers")

    args = parser.parse_args()
    return run_check(args) if args.command == "check" else run_build(args)


if __name__ == "__main__":
    sys.exit(main() or 0)
