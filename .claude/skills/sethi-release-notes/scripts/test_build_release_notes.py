#!/usr/bin/env python3
"""Offline tests for the release-notes renderers. No network, no Jira token.

Run: python3 .claude/skills/sethi-release-notes/scripts/test_build_release_notes.py
"""
from __future__ import annotations

import html.parser
import os
import re
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent))

import build_release_notes as brn
from openpyxl import load_workbook
from jira_api import JiraApiError, escape_jql, resolve_credentials, search_all, whoami

BASE = "https://team-anervea.atlassian.net"
TODAY = date(2026, 9, 30)

# The constructs Outlook's Word engine ignores. If any reaches the Outlook variant, the paste
# silently loses its formatting — the one failure this skill exists to prevent.
OUTLOOK_FORBIDDEN = r"var\(--|color-mix|clamp\(|env\(|position:\s*sticky|@media|googleapis|<style"

ISSUES = [
    {"key": "STH-673", "fields": {
        "summary": "Gate the audio fork <b>on</b> caption demand & GPU",
        "status": {"name": "Deployment"},
        "customfield_10020": [{"name": "STH Sprint 11"}]}},
    {"key": "STH-1420", "fields": {
        "summary": "POST /actions: create an action",
        "status": {"name": "Deployment"},
        "customfield_10020": [{"name": "STH Sprint 11"}]}},
    {"key": "STH-97", "fields": {
        "summary": "Numeric ordering check", "status": None, "customfield_10020": None}},
]


def make_sheet(issues=None, sender="Axe", has_workbook=True, scope=None):
    issues = ISSUES if issues is None else issues
    return brn.build_sheet(brn.to_rows(issues, BASE), "SETHI", "v0.0.3",
                           scope or brn.Scope("Deployment"), brn.sprint_names(issues), TODAY,
                           sender, has_workbook)


class TestRows(unittest.TestCase):
    def test_ordered_by_number_not_string(self):
        keys = [row.key for row in brn.to_rows(ISSUES, BASE)]
        self.assertEqual(keys, ["STH-97", "STH-673", "STH-1420"])

    def test_missing_status_does_not_crash(self):
        rows = {row.key: row for row in brn.to_rows(ISSUES, BASE)}
        self.assertEqual(rows["STH-97"].status, brn.UNKNOWN_STATUS)

    def test_url_points_at_the_browse_page(self):
        self.assertEqual(brn.to_rows(ISSUES, BASE)[0].url, f"{BASE}/browse/STH-97")

    def test_sprints_deduped_and_none_safe(self):
        self.assertEqual(brn.sprint_names(ISSUES), ["STH Sprint 11"])


class TestSheet(unittest.TestCase):
    def test_status_column_is_always_present(self):
        self.assertIn(brn.Kind.STATUS, [c.kind for c in make_sheet().columns])

    def test_band_counts_the_rows(self):
        self.assertEqual(make_sheet().band, "SETHI — 3 tickets")

    def test_both_renderers_read_the_same_column_set(self):
        """The DRY seam: column identity lives in the sheet, not in either renderer."""
        sheet = make_sheet()
        for column in sheet.columns:
            self.assertIn(column.header, brn.render_web(sheet))
            self.assertIn(column.header, brn.render_outlook(sheet))


class TestRendering(unittest.TestCase):
    def setUp(self):
        self.sheet = make_sheet()
        self.docs = {"web": brn.render_web(self.sheet),
                     "outlook": brn.render_outlook(self.sheet)}

    def test_both_parse(self):
        for name, doc in self.docs.items():
            with self.subTest(name):
                html.parser.HTMLParser().feed(doc)

    def test_summary_markup_is_escaped(self):
        for name, doc in self.docs.items():
            with self.subTest(name):
                self.assertIn("&lt;b&gt;", doc)
                self.assertNotIn("<b>on</b>", doc)

    def test_every_ticket_is_linked_once(self):
        for name, doc in self.docs.items():
            with self.subTest(name):
                self.assertEqual(doc.count("/browse/STH-"), len(ISSUES))

    def test_outlook_avoids_every_construct_word_ignores(self):
        hits = set(re.findall(OUTLOOK_FORBIDDEN, self.docs["outlook"]))
        self.assertEqual(hits, set(), f"Outlook variant carries {hits}")

    def test_web_defines_each_theme_token_in_all_three_places(self):
        """Bare :root, the prefers-color-scheme block, and the explicit [data-theme] block."""
        for token in ("--ink", "--page", "--link", "--band", "--sheet"):
            with self.subTest(token):
                self.assertGreaterEqual(self.docs["web"].count(f"{token}:"), 3)

    def test_web_body_sets_an_explicit_background(self):
        self.assertIn("body { margin:0; background:var(--page)", self.docs["web"])


class TestFilenameSafety(unittest.TestCase):
    def test_slug_strips_path_traversal(self):
        self.assertEqual(brn.slug("../../etc/passwd"), "etc-passwd")

    def test_every_version_lands_in_its_own_folder(self):
        self.assertEqual(brn.output_dir("SETHI", "v0.0.4", None),
                         brn.Path("sethi_release_notes/v0-0-4"))

    def test_two_versions_never_share_a_folder(self):
        self.assertNotEqual(brn.output_dir("SETHI", "v0.0.3", None),
                            brn.output_dir("SETHI", "v0.0.4", None))

    def test_a_version_cannot_walk_out_of_the_output_folder(self):
        self.assertEqual(brn.output_dir("SETHI", "../../etc", None),
                         brn.Path("sethi_release_notes/etc"))

    def test_an_explicit_out_wins_whole(self):
        self.assertEqual(brn.output_dir("SETHI", "v0.0.4", "/tmp/sheets"),
                         brn.Path("/tmp/sheets"))

    def test_slug_never_returns_empty(self):
        self.assertEqual(brn.slug("///"), "sheet")


class TestJqlSafety(unittest.TestCase):
    def test_quotes_in_a_status_cannot_break_out_of_the_literal(self):
        self.assertEqual(escape_jql('Code "Review"'), 'Code \\"Review\\"')

    def test_backslashes_are_escaped_before_quotes(self):
        self.assertEqual(escape_jql('a\\b'), 'a\\\\b')
        self.assertEqual(escape_jql('a\\"b'), 'a\\\\\\"b')


class TestCensus(unittest.TestCase):
    def test_a_truncated_read_is_refused_not_returned(self):
        class Truncated:
            def post(self, path, body):
                return {"issues": [{"key": "A-1"}], "isLast": False}

        with self.assertRaises(JiraApiError):
            search_all(Truncated(), "project = X", ["summary"])

    def test_pages_to_exhaustion(self):
        class Paged:
            def __init__(self):
                self.calls = 0

            def post(self, path, body):
                self.calls += 1
                if self.calls == 1:
                    return {"issues": [{"key": "A-1"}], "isLast": False, "nextPageToken": "t"}
                return {"issues": [{"key": "A-2"}], "isLast": True}

        self.assertEqual([i["key"] for i in search_all(Paged(), "p = X", ["summary"])],
                         ["A-1", "A-2"])



class TestCredentialResolution(unittest.TestCase):
    """A pair is taken WHOLE from one source — mixing sources names an account nobody holds."""

    def setUp(self):
        handle = tempfile.NamedTemporaryFile("w", suffix=".env", delete=False, encoding="utf-8")
        handle.write("JIRA_BASE_URL=https://file.example\n"
                     "JIRA_EMAIL=file@example.com\n"
                     "JIRA_API_TOKEN=file-token\n")
        handle.close()
        self.env_file = Path(handle.name)
        self.addCleanup(self.env_file.unlink)

    @mock.patch.dict(os.environ, {"ATLASSIAN_EMAIL": "atl@example.com",
                                  "ATLASSIAN_API_TOKEN": "atl-token",
                                  "JIRA_BASE_URL": "https://env.example"}, clear=True)
    def test_atlassian_pair_wins_over_the_file(self):
        env = resolve_credentials(self.env_file)
        self.assertEqual(env["JIRA_EMAIL"], "atl@example.com")
        self.assertEqual(env["JIRA_API_TOKEN"], "atl-token")

    @mock.patch.dict(os.environ, {"JIRA_EMAIL": "jira@example.com",
                                  "JIRA_API_TOKEN": "jira-token"}, clear=True)
    def test_jira_pair_is_used_and_borrows_the_files_base_url(self):
        env = resolve_credentials(self.env_file)
        self.assertEqual(env["JIRA_EMAIL"], "jira@example.com")
        self.assertEqual(env["JIRA_BASE_URL"], "https://file.example")

    @mock.patch.dict(os.environ, {"ATLASSIAN_EMAIL": "atl@example.com"}, clear=True)
    def test_an_incomplete_pair_is_ignored_not_half_used(self):
        """The email alone must not combine with the file's token."""
        env = resolve_credentials(self.env_file)
        self.assertEqual(env["JIRA_EMAIL"], "file@example.com")
        self.assertEqual(env["JIRA_API_TOKEN"], "file-token")

    @mock.patch.dict(os.environ, {}, clear=True)
    def test_falls_back_to_the_file_when_the_environment_is_empty(self):
        self.assertEqual(resolve_credentials(self.env_file)["JIRA_EMAIL"], "file@example.com")


class TestIdentityCheck(unittest.TestCase):
    def test_a_run_that_cannot_name_its_account_refuses_to_continue(self):
        class Anonymous:
            def get(self, path, params=None):
                return {}

        with self.assertRaises(JiraApiError):
            whoami(Anonymous())

    def test_the_display_name_is_returned_when_present(self):
        class Named:
            def get(self, path, params=None):
                return {"displayName": "rhidzkhan ahmad"}

        self.assertEqual(whoami(Named()), "rhidzkhan ahmad")


class TestCoveringMessage(unittest.TestCase):
    """The note ships with every sheet, so every fact in it is drawn, never typed."""

    def setUp(self):
        self.message = "\n".join(make_sheet().message)

    def test_it_names_the_client_the_version_and_the_count(self):
        self.assertIn("SETHI v0.0.3", self.message)
        self.assertIn("3 tickets", self.message)

    def test_it_names_the_status_actually_queried(self):
        self.assertIn("Deployment column", self.message)

    def test_it_names_both_sign_off_columns_as_rendered(self):
        self.assertIn(brn.SIGNOFF_HEADER, self.message)
        self.assertIn(brn.DECISION_HEADER, self.message)

    def test_it_offers_every_decision_the_dropdown_accepts(self):
        for choice in brn.DECISION_CHOICES:
            self.assertIn(choice, self.message)

    def test_it_is_signed_by_the_sender(self):
        self.assertEqual(make_sheet(sender="Marnie").message[-1], "Thanks,\nMarnie")

    def test_the_decision_header_is_quoted_without_doubled_punctuation(self):
        self.assertIn(f'under "{brn.DECISION_HEADER}"', "\n".join(make_sheet().message))
        self.assertNotIn("?.", "\n".join(make_sheet().message))

    def test_the_sign_off_newline_becomes_a_break_in_both_html_variants(self):
        for name, doc in (("web", brn.render_web(make_sheet())),
                          ("outlook", brn.render_outlook(make_sheet()))):
            with self.subTest(name):
                self.assertIn("Thanks,<br>Axe", doc)

    def test_it_promises_a_spreadsheet_only_when_one_is_written(self):
        self.assertIn("attached spreadsheet", "\n".join(make_sheet(has_workbook=True).message))
        self.assertNotIn("attached spreadsheet",
                         "\n".join(make_sheet(has_workbook=False).message))

    def test_a_single_ticket_is_not_pluralised(self):
        one = make_sheet(issues=ISSUES[:1])
        self.assertIn("the 1 ticket currently", "\n".join(one.message))

    def test_both_html_variants_carry_it(self):
        for name, doc in (("web", brn.render_web(make_sheet())),
                          ("outlook", brn.render_outlook(make_sheet()))):
            with self.subTest(name):
                self.assertIn("Please find below the release notes", doc)
                self.assertIn("Axe", doc)


class TestSender(unittest.TestCase):
    def test_a_hyphenated_git_handle_yields_the_name_a_reader_knows(self):
        self.assertEqual(brn.sender_from("Alfa-Axe", "rhidzkhan ahmad"), "Axe")

    def test_it_falls_back_to_the_jira_display_name(self):
        self.assertEqual(brn.sender_from(None, "rhidzkhan ahmad"), "Rhidzkhan")

    def test_a_blank_git_name_is_not_mistaken_for_a_name(self):
        self.assertEqual(brn.sender_from("   ", "rhidzkhan ahmad"), "Rhidzkhan")

    def test_a_plain_name_survives_whole(self):
        self.assertEqual(brn.sender_from("Marnie", "ignored"), "Marnie")


class TestSignOffColumns(unittest.TestCase):
    def test_both_sign_off_columns_are_present(self):
        kinds = [column.kind for column in make_sheet().columns]
        self.assertIn(brn.Kind.SIGNOFF, kinds)
        self.assertIn(brn.Kind.DECISION, kinds)

    def test_they_ship_empty_for_the_reviewer_to_fill(self):
        sheet = make_sheet()
        fillable = [c for c in sheet.columns
                    if c.kind in (brn.Kind.SIGNOFF, brn.Kind.DECISION)]
        for column in fillable:
            for row in sheet.rows:
                with self.subTest(column.header, key=row.key):
                    self.assertEqual(column.of(row).text, "")

    def test_the_dropdown_follows_the_column_rather_than_an_index(self):
        sheet = make_sheet()
        found = brn.decision_column(sheet)
        self.assertEqual(sheet.columns[found - 1].kind, brn.Kind.DECISION)


class TestWorkbook(unittest.TestCase):
    def setUp(self):
        self.path = Path(tempfile.mkdtemp()) / "sheet.xlsx"
        brn.render_xlsx(make_sheet(), self.path)
        self.book = load_workbook(self.path)

    def test_it_writes_a_table_sheet_and_a_covering_note_sheet(self):
        self.assertIn(brn.XL_NOTE_TAB, self.book.sheetnames)
        self.assertEqual(len(self.book.sheetnames), 2)

    def test_every_row_reaches_the_workbook(self):
        table = self.book.worksheets[0]
        self.assertEqual(table.max_row, len(make_sheet().rows) + 1)

    def test_the_ticket_cell_links_to_jira(self):
        table = self.book.worksheets[0]
        self.assertTrue(table.cell(row=2, column=2).hyperlink.target.endswith("STH-97"))

    def test_the_decision_column_carries_the_dropdown(self):
        table = self.book.worksheets[0]
        formulas = [v.formula1 for v in table.data_validations.dataValidation]
        self.assertEqual(formulas, ['"Yes,No,On Hold"'])

    def test_the_header_row_is_frozen_so_it_survives_scrolling(self):
        self.assertEqual(self.book.worksheets[0].freeze_panes, "A2")

    def test_a_tab_name_excel_would_reject_is_made_safe(self):
        self.assertEqual(brn.tab_name("a/b:c*d?e[f]g"), "a-b-c-d-e-f-g")
        self.assertLessEqual(len(brn.tab_name("x" * 60)), 31)


class TestCssContract(unittest.TestCase):
    def test_every_kind_value_has_a_matching_selector_in_the_web_css(self):
        """Renaming a Kind would otherwise silently orphan the stylesheet's selectors."""
        for kind in brn.Kind:
            with self.subTest(kind.value):
                self.assertIn(f".{kind.value} {{", brn.WEB_CSS)

    def test_every_kind_has_an_outlook_style(self):
        for kind in brn.Kind:
            with self.subTest(kind.value):
                self.assertIn(kind, brn.OL_CELL_STYLE)


class TestWording(unittest.TestCase):
    def test_a_single_ticket_is_not_pluralised(self):
        sheet = make_sheet([ISSUES[0]])
        self.assertEqual(sheet.band, "SETHI — 1 ticket")

    def test_the_meta_names_the_status_actually_queried(self):
        self.assertTrue(make_sheet().meta.startswith("Deployment column"))


class TestScopeWording(unittest.TestCase):
    """A scoped release and a whole column must never describe themselves the same way."""

    def test_a_whole_column_run_says_column_in_both_places(self):
        """The unscoped wording is unchanged by the arrival of --keys."""
        sheet = make_sheet(scope=brn.Scope("Deployment"))
        self.assertTrue(sheet.meta.startswith("Deployment column"))
        self.assertIn("currently in the Deployment column", "\n".join(sheet.message))

    def test_a_scoped_run_does_not_claim_the_whole_column(self):
        """A named set is not the column it sits in, and must not read as though it were."""
        sheet = make_sheet(scope=brn.Scope("Done", selected=True))
        self.assertIn("selected for this release", sheet.meta)
        message = "\n".join(sheet.message)
        self.assertIn("selected for this release", message)
        self.assertNotIn("currently in the Done column for this sprint", message)

    def test_no_wording_claims_a_ticket_was_released(self):
        """Membership comes from the operator's key list, which Jira cannot corroborate.

        The assertion names the bare word, not a phrase: a phrase no branch can emit would
        make this test unfailable, which is worse than not having it.
        """
        for scope in (brn.Scope("Deployment"), brn.Scope("Done", selected=True)):
            with self.subTest(selected=scope.selected):
                self.assertNotIn("released", "\n".join(make_sheet(scope=scope).message))

    def test_the_masthead_and_the_note_name_the_same_status(self):
        """The two descriptions of one sheet may never disagree about its status."""
        for scope in (brn.Scope("Deployment"), brn.Scope("Done", selected=True)):
            with self.subTest(scope.status):
                sheet = make_sheet(scope=scope)
                self.assertIn(scope.status, sheet.meta)
                self.assertIn(scope.status, "\n".join(sheet.message))


class TestStatusOf(unittest.TestCase):
    """One rule for reading a status, shared by the rows and the scope guard."""

    def test_it_reads_the_status_name(self):
        """The ordinary case."""
        self.assertEqual(brn.status_of({"fields": {"status": {"name": "Done"}}}), "Done")

    def test_every_shape_of_absence_becomes_the_same_unknown(self):
        """A missing fields, status or name must not produce three different answers."""
        for shape in ({}, {"fields": None}, {"fields": {}}, {"fields": {"status": None}},
                      {"fields": {"status": {}}}, {"fields": {"status": {"name": ""}}}):
            with self.subTest(repr(shape)):
                self.assertEqual(brn.status_of(shape), brn.UNKNOWN_STATUS)

    def test_the_rows_and_the_guard_read_a_status_identically(self):
        """Drift here would let the sheet display one status while the gate verified another."""
        issue = {"key": "STH-1", "fields": {"status": {"name": "QA Testing"}}}
        self.assertEqual(brn.to_rows([issue], BASE)[0].status, brn.status_of(issue))


class TestReleaseJql(unittest.TestCase):
    """The branch that decides both the query and the masthead wording."""

    def test_an_unscoped_run_queries_the_whole_status_column(self):
        """No keys means the original whole-column behaviour, unchanged."""
        jql = brn.release_jql("STH", "Deployment", ())
        self.assertEqual(jql, 'project = STH AND status = "Deployment" ORDER BY key ASC')

    def test_a_scoped_run_queries_the_keys_and_not_the_status(self):
        """The status is the masthead's claim, proven separately; it must not narrow the query."""
        jql = brn.release_jql("STH", "Done", ("STH-1", "STH-2"))
        self.assertEqual(jql, "project = STH AND key IN (STH-1, STH-2) ORDER BY key ASC")
        self.assertNotIn("status", jql)

    def test_the_project_is_escaped_in_both_modes(self):
        """project has no shape to validate against, so it is escaped rather than checked."""
        for keys in ((), ("STH-1",)):
            with self.subTest(scoped=bool(keys)):
                self.assertIn(escape_jql('a"b'), brn.release_jql('a"b', "Done", keys))


class TestKeyParsing(unittest.TestCase):
    """--keys is the one place operator text reaches the JQL, so it is validated, not escaped."""

    def test_absent_flag_is_an_empty_scope(self):
        """No flag and an empty flag both mean the whole column."""
        self.assertEqual(brn.parse_keys(None), ())
        self.assertEqual(brn.parse_keys(""), ())

    def test_commas_and_whitespace_both_separate(self):
        """Keys get pasted out of Jira and out of chat, which separate them differently."""
        self.assertEqual(brn.parse_keys("STH-1, STH-2\nSTH-3"), ("STH-1", "STH-2", "STH-3"))

    def test_keys_are_upper_cased_and_de_duplicated_in_order(self):
        """A duplicate would otherwise double a row and inflate the count in the covering note."""
        self.assertEqual(brn.parse_keys("sth-9,STH-9,sth-4"), ("STH-9", "STH-4"))

    def test_anything_that_is_not_a_key_is_refused(self):
        """The flag must not become the raw-JQL back door the skill deliberately lacks."""
        for hostile in ('STH-1") OR (project', "STH-1 ORDER BY key", "*", "STH-", "1-STH"):
            with self.subTest(hostile), self.assertRaises(ValueError):
                brn.parse_keys(hostile)


class TestVerifyScope(unittest.TestCase):
    """The two proofs that let a scoped run keep an honest masthead."""

    @staticmethod
    def issue(key, status="Done"):
        """A minimal issue carrying just the key and status the guard reads."""
        return {"key": key, "fields": {"status": {"name": status}}}

    def test_a_fully_present_matching_set_passes(self):
        """The happy path raises nothing."""
        brn.verify_scope([self.issue("STH-1"), self.issue("STH-2")], ("STH-1", "STH-2"), "Done")

    def test_a_key_that_did_not_come_back_is_refused(self):
        """A silently dropped key would shrink the release without changing its heading."""
        with self.assertRaises(JiraApiError) as caught:
            brn.verify_scope([self.issue("STH-1")], ("STH-1", "STH-2"), "Done")
        self.assertIn("STH-2", str(caught.exception))

    def test_a_ticket_in_another_status_is_refused(self):
        """Otherwise the masthead would claim a status the ticket does not hold."""
        with self.assertRaises(JiraApiError) as caught:
            brn.verify_scope([self.issue("STH-1", "QA Testing")], ("STH-1",), "Done")
        self.assertIn("QA Testing", str(caught.exception))

    def test_a_missing_status_object_is_refused_rather_than_defaulted(self):
        """An unreadable status is not evidence that the ticket holds the stated one."""
        with self.assertRaises(JiraApiError):
            brn.verify_scope([{"key": "STH-1", "fields": {}}], ("STH-1",), "Done")


class TestWriteVariants(unittest.TestCase):
    """The write stage, lifted out of main() so the variant filter is testable."""

    def test_only_the_requested_variant_is_written(self):
        """--variant web must not leave an .xlsx or an Outlook file beside it."""
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "v0-0-9"
            brn.write_variants(make_sheet(), "web", out, "sheet")
            self.assertEqual([f.name for f in out.iterdir()], ["sheet.html"])

    def test_all_writes_every_variant(self):
        """The default variant produces the full three-file package."""
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "v0-0-9"
            brn.write_variants(make_sheet(), brn.ALL_VARIANTS, out, "sheet")
            self.assertEqual(len(list(out.iterdir())), len(brn.VARIANTS))

    def test_it_creates_the_version_folder_it_was_given(self):
        """Every version gets its own folder, so one release never lands on another's."""
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "nested" / "v0-0-9"
            brn.write_variants(make_sheet(), "web", out, "sheet")
            self.assertTrue(out.is_dir())


if __name__ == "__main__":
    unittest.main(verbosity=2)
