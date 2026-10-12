#!/usr/bin/env python3
"""Render a Jira status column as a standalone release-notes sheet.

Three variants from one census:

  web      full sheet, light + dark themes, Google-hosted fonts — for viewing and sharing
  outlook  inline styles only, light only, table layout — survives a paste into Outlook
  xlsx     the working copy reviewers fill in — Yes/No/On Hold dropdown, colour rules

Jira access is delegated whole to jira-scripts/scripts/jira_api.py, which already owns auth,
retries, error shape, JQL escaping and the proven-exhaustive search. That file exists because
four copies of a Jira client had already drifted apart; this skill consumes it rather than
becoming the fifth. Only rendering and the CLI live here.

Stdlib plus openpyxl, which the workbook needs and which three sibling skills in this repo
already depend on. The HTML variants themselves use nothing outside the stdlib.
"""
from __future__ import annotations

import argparse
import html
import re
import subprocess
import sys
from dataclasses import dataclass
from datetime import date
from enum import Enum
from pathlib import Path
from typing import Callable

SKILLS_DIR = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(SKILLS_DIR / "jira-scripts" / "scripts"))

# Imported after the path insert above, which is what makes the sibling resolvable.
from jira_api import (
    Jira,
    JiraApiError,
    escape_jql,
    resolve_credentials,
    search_all,
    whoami,
)
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DEFAULT_STATUS = "Deployment"
DEFAULT_VERSION = "v0.0.1"
UNKNOWN_STATUS = "Unknown"
SLUG_FALLBACK = "sheet"
SPRINT_FIELD = "customfield_10020"
FIELDS = ["summary", "status", SPRINT_FIELD]

SIGNOFF_HEADER = "Business Team Sign Off"
DECISION_HEADER = "Can this be deployed?"
DECISION_CHOICES = ("Yes", "No", "On Hold")
ENVIRONMENT = "UAT"

SANS = "'Segoe UI', Calibri, Arial, sans-serif"
SERIF = "Georgia, 'Times New Roman', serif"
MONO = "Consolas, 'Courier New', monospace"


# ---------------------------------------------------------------- the view model
# Both renderers read this and nothing else. Every piece of CONTENT — the headings, the
# column set, their order, the band and footer wording — is settled once here, so adding or
# reordering a column is one edit rather than two that must be kept in step. Only styling is
# duplicated between the renderers, because their rendering contracts are incompatible.


class Kind(Enum):
    """What a column holds. Each renderer maps this to its own styling."""

    CLIENT = "client"
    TICKET = "ticket"
    DETAILS = "details"
    STATUS = "status"
    SIGNOFF = "signoff"
    DECISION = "decision"


@dataclass(frozen=True)
class Row:
    key: str
    url: str
    summary: str
    status: str


@dataclass(frozen=True)
class Cell:
    text: str
    href: str | None = None


@dataclass(frozen=True)
class Column:
    header: str
    kind: Kind
    of: Callable[[Row], Cell]


@dataclass(frozen=True)
class Sheet:
    title: str
    meta: str
    band: str
    message: tuple[str, ...]
    footer: str
    columns: tuple[Column, ...]
    rows: tuple[Row, ...]


def tally(count: int) -> str:
    """'91 tickets', '1 ticket'. One place, because the band and the message both say it."""
    return f"{count} ticket{'' if count == 1 else 's'}"


def covering_message(client: str, version: str, status: str, count: int, sender: str,
                     has_workbook: bool) -> tuple[str, ...]:
    """The note that goes out with the sheet, as paragraphs.

    It is written here rather than typed per release because every fact in it — the version,
    the count, the column — is already settled, and a covering note that contradicts the table
    under it is worse than no note. has_workbook keeps it honest: it may only promise the
    spreadsheet on a run that actually writes one.
    """
    fill_in = (f'Kindly review and complete the {SIGNOFF_HEADER} column, and mark each ticket '
               f'{" / ".join(DECISION_CHOICES)} under "{DECISION_HEADER}"')
    if has_workbook:
        fill_in += (" The attached spreadsheet has dropdowns in that column, so it is easiest "
                    "to fill it in there and send it back.")
    return (
        "Hi team,",
        f"Please find below the release notes for {client} {version}, covering the "
        f"{tally(count)} currently in the {status} column for this sprint. These have been "
        f"deployed to {ENVIRONMENT} and are ready for your review.",
        fill_in,
        "We'll proceed with the production deployment once all items are signed off. Let me "
        "know if you need any clarification on a particular ticket.",
        f"Thanks,\n{sender}",
    )


def build_sheet(rows: tuple[Row, ...], client: str, version: str, status: str,
                sprints: list[str], today: date, sender: str,
                has_workbook: bool) -> Sheet:
    """Settle every string and every column once, for all three renderers.

    The Status column is always present. The masthead names the column being released, but a
    heading is not a per-row guarantee: these sheets get filtered, re-sorted and pasted into
    mail, where a heading parts company with the rows underneath it.

    The two sign-off columns ship EMPTY on purpose — they are what the reviewer fills in, and
    the sheet must never arrive looking as though someone already answered for them.
    """
    meta_parts = [f"{status} column", *sprints, f"{today:%d %B %Y}"]
    return Sheet(
        title=f"{client} Release Notes — {version}",
        meta=" · ".join(meta_parts),
        band=f"{client} — {tally(len(rows))}",
        message=covering_message(client, version, status, len(rows), sender, has_workbook),
        footer=f"Pulled live from Jira on {today:%d %B %Y}.",
        columns=(
            Column("Client", Kind.CLIENT, lambda _row: Cell(client)),
            Column("Ticket", Kind.TICKET, lambda row: Cell(row.key, href=row.url)),
            Column("Details", Kind.DETAILS, lambda row: Cell(row.summary)),
            Column("Status", Kind.STATUS, lambda row: Cell(row.status)),
            Column(SIGNOFF_HEADER, Kind.SIGNOFF, lambda _row: Cell("")),
            Column(DECISION_HEADER, Kind.DECISION, lambda _row: Cell("")),
        ),
        rows=rows,
    )


# ---------------------------------------------------------------- Jira -> rows


def to_rows(issues: list[dict], base_url: str) -> tuple[Row, ...]:
    """Issues as rows, ordered by ticket number rather than by string."""
    rows = []
    for issue in issues:
        fields = issue.get("fields") or {}
        rows.append(Row(
            key=issue["key"],
            url=f"{base_url}/browse/{issue['key']}",
            summary=fields.get("summary") or "",
            status=(fields.get("status") or {}).get("name") or UNKNOWN_STATUS,
        ))
    return tuple(sorted(rows, key=lambda row: _number(row.key)))


def _number(key: str) -> tuple[str, int]:
    prefix, _, number = key.rpartition("-")
    return (prefix, int(number)) if number.isdigit() else (prefix, 0)


def sprint_names(issues: list[dict]) -> list[str]:
    """Every sprint named on the issues, first-seen order, deduped in one pass."""
    seen = {
        sprint["name"]: None
        for issue in issues
        for sprint in ((issue.get("fields") or {}).get(SPRINT_FIELD) or [])
        if sprint.get("name")
    }
    return list(seen)


# ---------------------------------------------------------------- Outlook variant
# Outlook on Windows renders mail with the Word engine, which ignores <style> blocks on
# paste, CSS custom properties, web fonts, color-mix(), clamp(), env(), @media and sticky
# positioning. Every rule here is therefore an inline style attribute, with bgcolor and the
# cellpadding/cellspacing attributes alongside, and fonts that ship with Windows.

OL_INK = "#2c221a"
OL_INK_SOFT = "#6d5c4d"
OL_RULE = "#e8dcd1"
OL_LINK = "#a8400f"
OL_ZEBRA = "#faf5f0"
OL_WHITE = "#ffffff"
OL_HEAD_BG = "#f1e6dc"
OL_HEAD_RULE = "#c9b6a4"
OL_HEAD_INK = "#4a3a2c"
OL_MAST_BG = "#3b2d22"
OL_MAST_INK = "#f7efe6"
OL_MAST_SUB = "#d8c6b4"
OL_MAST_RULE = "#5a483a"
OL_BAND_BG = "#9c4418"
OL_BAND_INK = "#fdf3ea"
OL_BORDER = "#ded0c3"
OL_FILLABLE = "#fffdf8"
OL_MSG_BG = "#fbf7f2"

OL_CELL = (f"padding:7px 12px; border-bottom:1px solid {OL_RULE}; font-family:{SANS}; "
           f"font-size:13px; line-height:1.45; color:{OL_INK}; vertical-align:top;")
OL_HEAD = (f"padding:8px 12px; border-bottom:2px solid {OL_HEAD_RULE}; "
           f"background-color:{OL_HEAD_BG}; font-family:{SANS}; font-size:11px; "
           f"font-weight:bold; color:{OL_HEAD_INK}; text-align:left; letter-spacing:0.5px; "
           "text-transform:uppercase;")

OL_CELL_STYLE = {
    Kind.CLIENT: f"{OL_CELL} white-space:nowrap; color:{OL_INK_SOFT}; font-size:12px;",
    Kind.TICKET: f"{OL_CELL} white-space:nowrap; font-family:{MONO}; font-size:12px;",
    Kind.DETAILS: OL_CELL,
    Kind.STATUS: f"{OL_CELL} white-space:nowrap; font-size:12px;",
    Kind.SIGNOFF: f"{OL_CELL} background-color:{OL_FILLABLE};",
    Kind.DECISION: f"{OL_CELL} background-color:{OL_FILLABLE};",
}
OL_HEAD_WIDTH = {Kind.CLIENT: "70px", Kind.TICKET: "90px", Kind.STATUS: "110px",
                 Kind.SIGNOFF: "150px", Kind.DECISION: "150px"}


def line_break(text: str) -> str:
    """Escaped text with its newlines as <br> — the sign-off is the only place one occurs."""
    return html.escape(text).replace("\n", "<br>")


def cell_content(column: Column, row: Row) -> tuple[str, str | None]:
    """A cell's escaped text, and its escaped href when the cell links.

    Shared because WHICH cell links and WHAT it says are content decisions, settled once
    in the Sheet. Only the anchor's attributes differ between the renderers, and those
    are styling, so each renderer still writes its own markup around this.
    """
    cell = column.of(row)
    return html.escape(cell.text), html.escape(cell.href) if cell.href else None


def render_outlook(sheet: Sheet) -> str:
    message = "\n".join(
        f'    <div style="padding-bottom:10px; font-family:{SANS}; font-size:13px; '
        f'line-height:1.5; color:{OL_INK};">{line_break(paragraph)}</div>'
        for paragraph in sheet.message
    )

    headers = "".join(
        f'<th style="{OL_HEAD}'
        + (f" width:{OL_HEAD_WIDTH[column.kind]};" if column.kind in OL_HEAD_WIDTH else "")
        + f'">{html.escape(column.header)}</th>'
        for column in sheet.columns
    )

    body = []
    for index, row in enumerate(sheet.rows):
        background = OL_ZEBRA if index % 2 else OL_WHITE
        cells = ""
        for column in sheet.columns:
            text, href = cell_content(column, row)
            if href:
                text = (f'<a href="{href}" '
                        f'style="color:{OL_LINK}; text-decoration:underline;">{text}</a>')
            cells += f'<td style="{OL_CELL_STYLE[column.kind]}">{text}</td>'
        body.append(f'        <tr bgcolor="{background}">{cells}</tr>')

    return f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>{html.escape(sheet.title)}</title>
</head>
<body style="margin:0; padding:16px; background-color:{OL_WHITE}; font-family:{SANS}; font-size:13px; color:{OL_INK};">
<table role="presentation" cellpadding="0" cellspacing="0" border="0" width="100%" style="border-collapse:collapse; width:100%; border:1px solid {OL_BORDER};">
  <tr><td bgcolor="{OL_MAST_BG}" style="padding:16px 18px; background-color:{OL_MAST_BG};">
    <div style="font-family:{SERIF}; font-size:20px; font-weight:bold; color:{OL_MAST_INK};">{html.escape(sheet.title)}</div>
    <div style="padding-top:8px; margin-top:8px; border-top:1px solid {OL_MAST_RULE}; font-family:{SANS}; font-size:12px; color:{OL_MAST_SUB};">{html.escape(sheet.meta)}</div>
  </td></tr>
  <tr><td bgcolor="{OL_BAND_BG}" style="padding:9px 18px; background-color:{OL_BAND_BG}; font-family:{SANS}; font-size:14px; font-weight:bold; color:{OL_BAND_INK};">{html.escape(sheet.band)}</td></tr>
  <tr><td bgcolor="{OL_MSG_BG}" style="padding:16px 18px 6px; background-color:{OL_MSG_BG}; border-bottom:1px solid {OL_BORDER};">
{message}
  </td></tr>
  <tr><td style="padding:0;">
    <table role="presentation" cellpadding="0" cellspacing="0" border="0" width="100%" style="border-collapse:collapse; width:100%;">
      <thead><tr bgcolor="{OL_HEAD_BG}">{headers}</tr></thead>
      <tbody>
{chr(10).join(body)}
      </tbody>
    </table>
  </td></tr>
</table>
<div style="margin-top:12px; font-family:{SANS}; font-size:11px; color:{OL_INK_SOFT};">{html.escape(sheet.footer)}</div>
</body>
</html>
"""


# ---------------------------------------------------------------- web variant

DARK_TOKENS = (
    "--page:#18120e; --sheet:#221a14; --rule:#44342a; --rule-soft:#33271f; "
    "--ink:#f0e6dc; --ink-soft:#b09a86; --masthead:#3a2c21; --mast-ink:#f7efe6; "
    "--mast-sub:#cdb8a4; --band:#8f3d13; --band-ink:#fdf3ea; --head:#32261e; "
    "--zebra:#271e17; --link:#f0904a; --fillable:#2a211a; --note:#1f1813; color-scheme:dark;"
)

WEB_CSS = (
    "  /* One full-bleed ledger: masthead, client band, single scrolling table. */\n"
    "  :root {\n"
    "    --page:#f6f1ec; --sheet:#ffffff; --rule:#ded0c3; --rule-soft:#ece1d7;\n"
    "    --ink:#2c221a; --ink-soft:#6d5c4d; --masthead:#3b2d22; --mast-ink:#f7efe6;\n"
    "    --mast-sub:#d8c6b4; --band:#9c4418; --band-ink:#fdf3ea; --head:#f1e6dc;\n"
    "    --zebra:#faf5f0; --link:#a8400f; --fillable:#fffdf8; --note:#fbf7f2;\n"
    '    --font-display:"Source Serif 4",Georgia,serif;\n'
    '    --font-body:"IBM Plex Sans",-apple-system,BlinkMacSystemFont,"Segoe UI",system-ui,sans-serif;\n'
    '    --font-mono:"IBM Plex Mono",ui-monospace,Menlo,monospace;\n'
    "    color-scheme: light;\n"
    "    padding-top: env(safe-area-inset-top, 0px);\n"
    "    padding-bottom: env(safe-area-inset-bottom, 0px);\n"
    "  }\n"
    '  @media (prefers-color-scheme: dark) { :root:not([data-theme="light"]) { ' + DARK_TOKENS + " } }\n"
    '  :root[data-theme="dark"] { ' + DARK_TOKENS + " }\n"
    "  body { margin:0; background:var(--page); color:var(--ink); font-family:var(--font-body); font-size:15px; line-height:1.5; }\n"
    "  img { max-width:100%; }\n"
    "  [hidden] { display:none !important; }\n"
    "  .wrap { max-width:1180px; margin:0 auto; padding-inline:16px; padding-block:24px 48px; }\n"
    "  .sheet { background:var(--sheet); border:1px solid var(--rule); border-radius:4px; overflow:hidden; }\n"
    "  .masthead { background:var(--masthead); color:var(--mast-ink); padding:18px 20px 14px; }\n"
    "  .masthead h1 { margin:0; font-family:var(--font-display); font-weight:700; font-size:clamp(1.2rem,1rem + 1.1vw,1.6rem); letter-spacing:-0.01em; text-wrap:balance; }\n"
    "  .masthead .meta { margin:8px 0 0; color:var(--mast-sub); font-size:0.83rem; border-top:1px solid rgb(255 255 255 / 0.14); padding-top:10px; }\n"
    "  .band { background:var(--band); color:var(--band-ink); padding:10px 20px; font-weight:600; font-size:0.95rem; }\n"
    "  .scroller { overflow-x:auto; }\n"
    "  table { width:100%; min-width:640px; border-collapse:collapse; font-variant-numeric:tabular-nums; }\n"
    "  th, td { text-align:left; vertical-align:top; padding:9px 14px; border-bottom:1px solid var(--rule-soft); }\n"
    "  thead th { position:sticky; top:env(safe-area-inset-top, 0px); background:var(--head); color:var(--ink); font-size:0.72rem; font-weight:600; text-transform:uppercase; letter-spacing:0.07em; border-bottom:1px solid var(--rule); white-space:nowrap; }\n"
    "  tbody tr:nth-child(even) { background:var(--zebra); }\n"
    "  tbody tr:last-child td { border-bottom:none; }\n"
    "  .client { width:1%; white-space:nowrap; color:var(--ink-soft); font-size:0.82rem; letter-spacing:0.04em; }\n"
    "  .ticket { width:1%; white-space:nowrap; font-family:var(--font-mono); font-size:0.86rem; }\n"
    "  .ticket a { color:var(--link); text-decoration:none; border-bottom:1px solid rgb(168 64 15 / 0.4); }\n"
    "  .ticket a:hover { border-bottom-color:var(--link); }\n"
    "  .ticket a:focus-visible { outline:2px solid var(--link); outline-offset:2px; border-radius:2px; }\n"
    "  .details { min-width:0; }\n"
    "  .status { width:1%; white-space:nowrap; color:var(--ink-soft); font-size:0.82rem; }\n"
    "  .signoff { width:150px; min-width:150px; background:var(--fillable); }\n"
    "  .decision { width:150px; min-width:150px; background:var(--fillable); }\n"
    "  .note { padding:18px 20px 8px; background:var(--note); border-bottom:1px solid var(--rule); }\n"
    "  .note p { margin:0 0 10px; max-width:68ch; }\n"
    "  .note p:last-child { margin-bottom:10px; }\n"
    "  .foot { margin:14px 2px 0; color:var(--ink-soft); font-size:0.8rem; }\n"
)


def render_web(sheet: Sheet) -> str:
    message = "\n".join(f"      <p>{line_break(paragraph)}</p>" for paragraph in sheet.message)

    headers = "".join(f'<th scope="col">{html.escape(c.header)}</th>' for c in sheet.columns)

    body = []
    for row in sheet.rows:
        cells = ""
        for column in sheet.columns:
            text, href = cell_content(column, row)
            if href:
                text = f'<a href="{href}" target="_blank" rel="noopener">{text}</a>'
            cells += f'<td class="{column.kind.value}">{text}</td>'
        body.append(f"          <tr>{cells}</tr>")

    return f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<title>{html.escape(sheet.title)}</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=IBM+Plex+Mono:wght@500&family=IBM+Plex+Sans:wght@400;500;600&family=Source+Serif+4:opsz,wght@8..60,700&display=swap">
<style>
{WEB_CSS}</style>
</head>
<body>
<div class="wrap">
  <div class="sheet">
    <header class="masthead">
      <h1>{html.escape(sheet.title)}</h1>
      <p class="meta">{html.escape(sheet.meta)}</p>
    </header>
    <div class="band">{html.escape(sheet.band)}</div>
    <section class="note" aria-label="Covering note">
{message}
    </section>
    <div class="scroller">
      <table>
        <thead><tr>{headers}</tr></thead>
        <tbody>
{chr(10).join(body)}
        </tbody>
      </table>
    </div>
  </div>
  <p class="foot">{html.escape(sheet.footer)}</p>
</div>
</body>
</html>
"""


# ---------------------------------------------------------------- xlsx variant
# The working copy. The HTML variants are read; this one is FILLED IN and sent back, so the
# decision column is a dropdown rather than free text — a sign-off sheet that accepts "yes pls"
# and "Y" in the same column cannot be counted without a human re-reading all of it.

XL_HEAD_FILL = "F1E6DC"
XL_HEAD_INK = "4A3A2C"
XL_RULE = "DED0C3"
XL_FILLABLE = "FFFDF8"
XL_DECISION_STYLES = {
    "Yes": ("C6EFCE", "006100"),
    "No": ("FFC7CE", "9C0006"),
    "On Hold": ("FFEB9C", "9C5700"),
}
XL_WIDTH = {
    Kind.CLIENT: 12, Kind.TICKET: 14, Kind.DETAILS: 78, Kind.STATUS: 16,
    Kind.SIGNOFF: 26, Kind.DECISION: 24,
}
XL_NOTE_TAB = "Covering note"
XL_TAB_FORBIDDEN = r"[\\/?*\[\]:]"
XL_TAB_LIMIT = 31


def tab_name(title: str) -> str:
    """Excel rejects a worksheet name over 31 chars or carrying []:*?/\\ — so neither reaches it."""
    return re.sub(XL_TAB_FORBIDDEN, "-", title)[:XL_TAB_LIMIT] or SLUG_FALLBACK


def decision_column(sheet: Sheet) -> int | None:
    """The 1-based index of the column the dropdown belongs on, or None if there is none.

    Found by Kind, never hardcoded: the dropdown and the colour rules must follow the column
    if build_sheet ever reorders them, and a validation left on a stale index is invisible.
    """
    for index, column in enumerate(sheet.columns, start=1):
        if column.kind is Kind.DECISION:
            return index
    return None


def render_xlsx(sheet: Sheet, path: Path) -> None:
    """Write the fill-in workbook: one table sheet, one covering-note sheet."""
    book = Workbook()
    table = book.active
    table.title = tab_name(sheet.title)

    rule = Side(style="thin", color=XL_RULE)
    border = Border(left=rule, right=rule, top=rule, bottom=rule)
    fillable = PatternFill("solid", fgColor=XL_FILLABLE)

    for index, column in enumerate(sheet.columns, start=1):
        cell = table.cell(row=1, column=index, value=column.header)
        cell.font = Font(bold=True, color=XL_HEAD_INK, size=10)
        cell.fill = PatternFill("solid", fgColor=XL_HEAD_FILL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = border
        table.column_dimensions[get_column_letter(index)].width = XL_WIDTH[column.kind]

    for offset, row in enumerate(sheet.rows):
        for index, column in enumerate(sheet.columns, start=1):
            content = column.of(row)
            cell = table.cell(row=2 + offset, column=index, value=content.text)
            cell.alignment = Alignment(vertical="top", wrap_text=column.kind is Kind.DETAILS)
            cell.border = border
            if content.href:
                cell.hyperlink = content.href
                cell.style = "Hyperlink"
            if column.kind in (Kind.SIGNOFF, Kind.DECISION):
                cell.fill = fillable

    last = len(sheet.rows) + 1
    index = decision_column(sheet)
    if index is not None and sheet.rows:
        letter = get_column_letter(index)
        span = f"{letter}2:{letter}{last}"
        choices = ",".join(DECISION_CHOICES)
        validation = DataValidation(type="list", formula1=f'"{choices}"', allow_blank=True,
                                    showDropDown=False, errorStyle="stop")
        validation.error = f"Choose one of: {', '.join(DECISION_CHOICES)}."
        validation.errorTitle = "Not a sign-off answer"
        table.add_data_validation(validation)
        validation.add(span)
        for answer, (fill, ink) in XL_DECISION_STYLES.items():
            table.conditional_formatting.add(span, CellIsRule(
                operator="equal", formula=[f'"{answer}"'],
                fill=PatternFill("solid", fgColor=fill), font=Font(color=ink, bold=True)))

    table.freeze_panes = "A2"
    table.auto_filter.ref = f"A1:{get_column_letter(len(sheet.columns))}{last}"
    table.print_title_rows = "1:1"
    table.page_setup.orientation = "landscape"
    table.page_setup.fitToWidth = 1
    table.page_setup.fitToHeight = 0
    table.sheet_properties.pageSetUpPr.fitToPage = True

    note = book.create_sheet(XL_NOTE_TAB)
    note.column_dimensions["A"].width = 110
    for offset, paragraph in enumerate(sheet.message, start=1):
        cell = note.cell(row=offset, column=1, value=paragraph)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        note.row_dimensions[offset].height = 15 * (1 + len(paragraph) // 100) + 6

    book.save(path)


# ---------------------------------------------------------------- entry point

def write_web(sheet: Sheet, path: Path) -> None:
    path.write_text(render_web(sheet), encoding="utf-8")


def write_outlook(sheet: Sheet, path: Path) -> None:
    path.write_text(render_outlook(sheet), encoding="utf-8")


ALL_VARIANTS = "all"
WORKBOOK = "xlsx"
VARIANTS = {
    "web": (write_web, ".html"),
    "outlook": (write_outlook, "-outlook.html"),
    WORKBOOK: (render_xlsx, ".xlsx"),
}


def sender_from(git_name: str | None, jira_name: str) -> str:
    """Who the covering note is signed by.

    The git identity first, because it is how the operator names themselves on this machine
    ("Alfa-Axe" -> "Axe"); the Jira display name as the fallback. --from overrides both, and
    exists because neither source is the name a business reader will recognise on every setup.
    """
    candidate = (git_name or "").strip() or jira_name.strip()
    tail = candidate.replace("_", "-").split("-")[-1].strip()
    return (tail or candidate).split()[0].title() if candidate else "the release team"


def git_user_name() -> str | None:
    """`git config user.name`, or None wherever git is absent or says nothing."""
    try:
        done = subprocess.run(["git", "config", "user.name"], capture_output=True,
                              text=True, timeout=5, check=False)
    except (OSError, subprocess.SubprocessError):
        return None
    return done.stdout.strip() or None


def slug(value: str) -> str:
    """A filename-safe token. User-supplied text reaches a path, so it is not trusted."""
    cleaned = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    return cleaned or SLUG_FALLBACK


def output_dir(client: str, version: str, override: str | None) -> Path:
    """Every version gets its own folder, so one release's sheets never mix with another's.

    Both segments go through slug(): --version reaches a directory name, and a raw
    "../.." would walk out of the output folder entirely.
    """
    if override:
        return Path(override)
    return Path(f"{slug(client)}_release_notes") / slug(version)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--project", required=True, help="Board key, e.g. STH")
    parser.add_argument("--status", default=DEFAULT_STATUS,
                        help=f'Status column (default: "{DEFAULT_STATUS}")')
    parser.add_argument("--client", help="Client name in the band (default: the project key)")
    parser.add_argument("--version", default=DEFAULT_VERSION,
                        help="Version label in the masthead")
    parser.add_argument("--variant", choices=[*VARIANTS, ALL_VARIANTS], default=ALL_VARIANTS)
    parser.add_argument("--from", dest="sender",
                        help="Name the covering note is signed by (default: your git user)")
    parser.add_argument("--out", help="Output folder (default: <client>_release_notes/<version>)")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    client = args.client or args.project
    today = date.today()

    try:
        jira = Jira(resolve_credentials())
        account = whoami(jira)
        print(f"acting as: {account}", file=sys.stderr)

        jql = (f'project = {escape_jql(args.project)} AND '
               f'status = "{escape_jql(args.status)}" ORDER BY key ASC')
        print(f"jql: {jql}", file=sys.stderr)
        issues = search_all(jira, jql, FIELDS)
    except JiraApiError as err:
        print(f"BLOCKED: {err}", file=sys.stderr)
        return 1

    if not issues:
        print("No issues matched — nothing written.", file=sys.stderr)
        return 1

    sender = args.sender or sender_from(git_user_name(), account)
    has_workbook = args.variant in (WORKBOOK, ALL_VARIANTS)
    sheet = build_sheet(to_rows(issues, jira.base_url), client, args.version, args.status,
                        sprint_names(issues), today, sender, has_workbook)

    out_dir = output_dir(client, args.version, args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = f"{slug(client)}-release-notes-{slug(args.version)}"

    for name, (write, suffix) in VARIANTS.items():
        if args.variant not in (name, ALL_VARIANTS):
            continue
        path = out_dir / f"{stem}{suffix}"
        write(sheet, path)
        print(f"wrote {path} ({path.stat().st_size:,} bytes, {len(sheet.rows)} rows)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
