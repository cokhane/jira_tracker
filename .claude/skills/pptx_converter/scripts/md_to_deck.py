#!/usr/bin/env python
"""Convert an All-Boards sprint-report Markdown file into a themed PPTX.

Unlike build_allboards_deck.py (which recomputes every number from the tracker DBs and a
pre-authored narrative JSON), this reads the *already-authored* Markdown report as the source
of truth and renders it slide-for-slide. That means whatever you edited/approved in the .md is
exactly what lands in the deck — no DB, no API key, no narrative files needed.

It parses the current AllBoards_Sprint_Report_<date>.md layout:
  - Overview — All Boards table
  - per board: Closing/Kickoff "Team Load & Delivery" funnel tables (9 columns),
    Key Achievements bullets, Roadblocks bullets
Styling reuses the Alfa Kinetic theme helpers in app.report_render.

    python scripts/md_to_deck.py [path/to/report.md]

With no argument it converts the newest AllBoards_Sprint_Report_*.md in ~/Documents/sprint_report/.
The .pptx is written next to the source .md (same basename).
"""
from __future__ import annotations

import glob
import os
import re
import sys

def _find_repo_root() -> str:
    """Walk up from this file until the jira_tracker root (holding app/report_render.py)."""
    here = os.path.dirname(os.path.abspath(__file__))
    while here != os.path.dirname(here):
        if os.path.isfile(os.path.join(here, "app", "report_render.py")):
            return here
        here = os.path.dirname(here)
    raise SystemExit("Could not locate jira_tracker root (app/report_render.py not found)")


REPO_ROOT = _find_repo_root()
sys.path.insert(0, REPO_ROOT)
os.chdir(REPO_ROOT)

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402

from app.report_render import (  # noqa: E402
    ACCENT, ACCENT_SOFT, CARD, CREAM, EMU_H, EMU_W, INK, MARGIN, MUTED, WHITE,
    _bg, _eyebrow_title, _fill, _footer, _stat_cards, _table, _text,
)

SUBTITLE = RGBColor(0xD8, 0xC8, 0xBA)
OUTPUT_DIR = os.path.expanduser("~/Documents/sprint_report")

# Shorter headers so the 9-column funnel table fits a 16:9 slide.
FUNNEL_HEADERS = ["Developer", "Assigned", "Completed", "Incomplete",
                  "Code Rev.", "QA Test", "Merged", "Other", "Delivered"]
FUNNEL_WIDTHS = [Inches(2.9)] + [Inches(1.129)] * 8


# ---------- markdown parsing ----------


def _clean(text: str) -> str:
    """Strip a markdown cell/inline: drop bold markers and surrounding whitespace."""
    return text.replace("**", "").strip()


def _parse_table(lines: list[str], i: int) -> tuple[list[str], list[list[str]], int]:
    """Parse a pipe table starting at line i (a '|' row). Returns headers, rows, next index."""
    def cells(row: str) -> list[str]:
        return [_clean(c) for c in row.strip().strip("|").split("|")]

    headers = cells(lines[i])
    i += 2  # skip header row and the |---| separator
    rows: list[list[str]] = []
    while i < len(lines) and lines[i].lstrip().startswith("|"):
        rows.append(cells(lines[i]))
        i += 1
    return headers, rows, i


def _bullets(lines: list[str], i: int) -> tuple[list[str], int]:
    """Collect consecutive '- ' bullet lines starting at/after i (skipping blanks)."""
    out: list[str] = []
    while i < len(lines):
        s = lines[i].strip()
        if s.startswith("- "):
            out.append(_clean(s[2:]))
        elif s and not s.startswith("- "):
            break
        i += 1
    return out, i


def parse_report(text: str) -> dict:
    """Parse the AllBoards markdown into a structured dict of slides."""
    lines = text.splitlines()
    doc: dict = {"subtitle": "", "overview": None, "boards": []}
    board: dict | None = None
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("_Generated") and not doc["subtitle"]:
            doc["subtitle"] = _clean(stripped.strip("_"))
            i += 1
            continue

        if stripped.startswith("## Overview"):
            i += 1
            while i < len(lines) and not lines[i].lstrip().startswith("|"):
                i += 1
            headers, rows, i = _parse_table(lines, i)
            doc["overview"] = {"headers": headers, "rows": rows}
            continue

        # A board starts at a top-level '# ' heading that is not the doc title.
        if stripped.startswith("# ") and " — " in stripped and "Closure" in stripped:
            board = {"name": stripped[2:].split(" — ")[0].strip(),
                     "heading": stripped[2:].strip(), "closing_line": "", "kickoff_line": "",
                     "tables": [], "achievements": [], "roadblocks": []}
            doc["boards"].append(board)
            i += 1
            continue

        if board is not None and stripped.startswith("- Closing sprint:"):
            board["closing_line"] = _clean(stripped[len("- "):])
            i += 1
            continue
        if board is not None and stripped.startswith("- Kickoff sprint:"):
            board["kickoff_line"] = _clean(stripped[len("- "):])
            i += 1
            continue

        if board is not None and stripped.startswith("## ") and "Team Load & Delivery" in stripped:
            title = stripped[3:].strip()
            i += 1
            while i < len(lines) and not lines[i].lstrip().startswith("|"):
                i += 1
            headers, rows, i = _parse_table(lines, i)
            notes, i = _bullets(lines, i)
            board["tables"].append({"title": title, "headers": headers, "rows": rows,
                                    "notes": notes})
            continue

        if board is not None and stripped.startswith("## ") and "Key Achievements" in stripped:
            bullets, i = _bullets(lines, i + 1)
            board["achievements"] = bullets
            continue

        if board is not None and stripped.startswith("## ") and stripped.endswith("Roadblocks"):
            bullets, i = _bullets(lines, i + 1)
            board["roadblocks"] = bullets
            continue

        i += 1
    return doc


def _cell(row: list[str], headers: list[str], name: str) -> str:
    for h, v in zip(headers, row):
        if name.lower() in h.lower():
            return v
    return ""


def _sprint_dates(line: str) -> str:
    """'Closing sprint: **NEX Sprint 4** (2026-07-13 → 2026-07-20)' -> the (...) part."""
    m = re.search(r"\((.*?)\)", line)
    return m.group(1) if m else ""


def _sprint_name(line: str) -> str:
    m = re.search(r":\s*(.*?)\s*\(", line)
    return m.group(1).strip() if m else ""


# ---------- slide builders ----------


def _slide(prs, color):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, color)
    return s


def title_slide(prs, doc):
    s = _slide(prs, INK)
    _text(s, MARGIN, Inches(0.85), EMU_W - 2 * MARGIN, Inches(0.5),
          [("ANERVEA", 14, ACCENT_SOFT, True)])
    _text(s, MARGIN, Inches(1.35), EMU_W - 2 * MARGIN, Inches(1.2),
          [("Engineering Sprint Report", 44, WHITE, True)])
    _text(s, MARGIN, Inches(2.55), EMU_W - 2 * MARGIN, Inches(0.4),
          [("ALL BOARDS  ·  SPRINT CLOSURE & KICKOFF", 14, ACCENT_SOFT, True)])
    ov = doc["overview"]
    total = next((r for r in ov["rows"] if "total" in r[0].lower()), None) if ov else None
    cards = []
    if total:
        cards = [
            ("Boards", len(doc["boards"]), "reporting"),
            ("Shipped", _cell(total, ov["headers"], "Shipped"), "this cycle"),
            ("Carry-over", _cell(total, ov["headers"], "Carry-over"), "into kickoff"),
            ("Committed", _cell(total, ov["headers"], "Committed"), "next sprint"),
        ]
        _stat_cards(s, cards, Inches(4.4), on_dark=True)
    _footer(s, doc["subtitle"].upper(), on_dark=True)


def overview_slide(prs, doc):
    ov = doc["overview"]
    if not ov:
        return
    s = _slide(prs, CREAM)
    _eyebrow_title(s, "Portfolio · all boards", "Sprint Overview")
    n = len(ov["headers"])
    first = Inches(2.3)
    rest = int((EMU_W - 2 * MARGIN - first) / (n - 1))
    _table(s, ov["headers"], ov["rows"], Inches(2.0),
           col_widths=[first] + [rest] * (n - 1), height=Inches(2.8))
    _footer(s, "CONFIDENTIAL · ENGINEERING LEADERSHIP")


def board_divider(prs, b):
    s = _slide(prs, INK)
    _text(s, MARGIN, Inches(1.0), EMU_W - 2 * MARGIN, Inches(0.5),
          [("ANERVEA", 14, ACCENT_SOFT, True)])
    _text(s, MARGIN, Inches(1.5), EMU_W - 2 * MARGIN, Inches(1.2),
          [(b["name"], 48, WHITE, True)])
    closing, kickoff = _sprint_name(b["closing_line"]), _sprint_name(b["kickoff_line"])
    _text(s, MARGIN, Inches(2.7), EMU_W - 2 * MARGIN, Inches(0.4),
          [(f"{closing} CLOSURE  ·  {kickoff} KICKOFF", 14, ACCENT_SOFT, True)])
    _text(s, MARGIN, Inches(3.15), EMU_W - 2 * MARGIN, Inches(0.8),
          [(f"{closing}:  {_sprint_dates(b['closing_line'])}", 13, SUBTITLE, False),
           (f"{kickoff}:  {_sprint_dates(b['kickoff_line'])}", 13, SUBTITLE, False)])
    _footer(s, f"CONFIDENTIAL · {closing}", on_dark=True)


def funnel_slide(prs, b, tbl):
    s = _slide(prs, CREAM)
    _eyebrow_title(s, "Team · load & delivery", tbl["title"])
    # Map every source row onto the shortened, fixed funnel header order.
    src = tbl["headers"]
    rows = []
    for r in tbl["rows"]:
        rows.append([
            _cell(r, src, "Developer") or r[0],
            _cell(r, src, "Assigned"), _cell(r, src, "Completed"), _cell(r, src, "Incomplete"),
            _cell(r, src, "Code Review"), _cell(r, src, "QA"), _cell(r, src, "Merged"),
            _cell(r, src, "Other"), _cell(r, src, "Delivered"),
        ])
    _table(s, FUNNEL_HEADERS, rows, Inches(1.7), col_widths=FUNNEL_WIDTHS, height=Inches(3.6))
    caption = tbl["notes"][0] if tbl["notes"] else ""
    if caption:
        _text(s, MARGIN, Inches(5.6), EMU_W - 2 * MARGIN, Inches(1.3),
              [(caption, 10, MUTED, False)])
    _footer(s, "CONFIDENTIAL · TEAM DELIVERY")


def achievements_slide(prs, b):
    if not b["achievements"]:
        return
    s = _slide(prs, CREAM)
    closing = _sprint_name(b["closing_line"])
    _eyebrow_title(s, "Sprint closure", f"{closing} — Key Achievements")
    runs = [(f"•  {a}", 14, INK, False) for a in b["achievements"]]
    _text(s, MARGIN, Inches(1.9), EMU_W - 2 * MARGIN, Inches(5.0), runs)
    _footer(s, f"CONFIDENTIAL · {closing}")


def roadblocks_slide(prs, b):
    if not b["roadblocks"]:
        return
    s = _slide(prs, CREAM)
    closing = _sprint_name(b["closing_line"])
    _eyebrow_title(s, "Risks & blockers", f"{closing} — Roadblocks & Next Steps")
    n = len(b["roadblocks"])
    top0, gap = Inches(1.9), Inches(0.2)
    avail = EMU_H - top0 - Inches(0.5)
    h = int((avail - gap * (n - 1)) / n)
    for idx, rb in enumerate(b["roadblocks"]):
        top = top0 + idx * (h + gap)
        rect = s.shapes.add_shape(1, MARGIN, top, EMU_W - 2 * MARGIN, h)
        _fill(rect, CARD)
        lead, _, rest = rb.partition(" — ")
        runs = [(lead.strip(), 14, ACCENT, True)]
        if rest.strip():
            runs.append((rest.strip(), 11.5, INK, False))
        _text(s, MARGIN + Inches(0.3), top + Inches(0.15), EMU_W - 2 * MARGIN - Inches(0.6),
              h - Inches(0.3), runs)
    _footer(s, f"CONFIDENTIAL · {closing}")


def closing_slide(prs, doc):
    s = _slide(prs, INK)
    _text(s, MARGIN, Inches(2.3), EMU_W - 2 * MARGIN, Inches(0.5),
          [("ANERVEA", 14, ACCENT_SOFT, True)])
    _text(s, MARGIN, Inches(2.8), EMU_W - 2 * MARGIN, Inches(1.0),
          [("Engineering Leadership", 40, WHITE, True)])
    _text(s, MARGIN, Inches(4.0), EMU_W - 2 * MARGIN, Inches(0.4),
          [("DISTRIBUTION   Engineering Team · Product Lead · CTO · QA · DevOps", 12,
            ACCENT_SOFT, True)])
    _footer(s, doc["subtitle"], on_dark=True)


def build_deck(doc) -> Presentation:
    prs = Presentation()
    prs.slide_width = EMU_W
    prs.slide_height = EMU_H
    title_slide(prs, doc)
    overview_slide(prs, doc)
    for b in doc["boards"]:
        board_divider(prs, b)
        for tbl in b["tables"]:
            funnel_slide(prs, b, tbl)
        achievements_slide(prs, b)
        roadblocks_slide(prs, b)
    closing_slide(prs, doc)
    return prs


def _default_md() -> str:
    matches = sorted(glob.glob(os.path.join(OUTPUT_DIR, "AllBoards_Sprint_Report_*.md")))
    if not matches:
        raise SystemExit(f"No AllBoards_Sprint_Report_*.md found in {OUTPUT_DIR}")
    return matches[-1]


def main():
    md_path = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else _default_md()
    with open(md_path, encoding="utf-8") as fh:
        doc = parse_report(fh.read())
    prs = build_deck(doc)
    out = os.path.splitext(md_path)[0] + ".pptx"
    prs.save(out)
    print(f"Parsed {len(doc['boards'])} boards -> {len(prs.slides._sldIdLst)} slides")
    print(f"Saved -> {out}")


if __name__ == "__main__":
    main()
